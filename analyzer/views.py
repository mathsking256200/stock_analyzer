import hashlib
import io
import json
import uuid

import razorpay
from django.conf import settings
from django.contrib import messages
from django.core.mail import EmailMessage, send_mail
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, legal
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Payment, Stock, User
from .utils import (
    WATCHLIST_BASE_HEADERS,
    WATCHLIST_FULL_HEADERS,
    build_watchlist_report_rows,
    get_candlestick_data,
    normalize_symbol,
    safe_float,
    scrape_peers,
    scrape_stock_summary,
    scrape_table,
    scrape_watchlist_with_selenium,
)


def analyze_stock(request):
    if request.method == "POST":
        slug = request.POST.get("stock_slug")

        if slug:
            return redirect("stock_detail", slug=slug)

    return redirect("index")



def index(request):
    sectors = Stock.objects.values_list("sector", flat=True).distinct()
    selected_sector = request.GET.get("sector", "")
    stocks = Stock.objects.all()

    if selected_sector:
        stocks = stocks.filter(sector=selected_sector)

    context = {
        "sectors": sectors,
        "stocks": stocks,
        "selected_sector": selected_sector,
    }

    return render(request, "index.html", context)



def stock_detail(request, slug):
    stock = get_object_or_404(Stock, slug=slug)
    period = request.GET.get("period", "1y")

    period_options = [
        ("1mo", "1M"),
        ("6mo", "6M"),
        ("1y", "1Y"),
        ("5y", "5Y"),
        ("max", "MAX"),
    ]

    clean_symbol = normalize_symbol(stock.symbol)

    summary = scrape_stock_summary(clean_symbol)

    pnl_headers, pnl = scrape_table(clean_symbol, "pnl")
    bs_headers, bs = scrape_table(clean_symbol, "bs")
    cf_headers, cf = scrape_table(clean_symbol, "cf")
    ratio_headers, ratios = scrape_table(clean_symbol, "ratios")
    share_headers, shareholding = scrape_table(clean_symbol, "shareholding")

    peer_headers, peer_rows = scrape_peers(clean_symbol)
    candles = get_candlestick_data(clean_symbol, period)

    context = {
        "stock": stock,
        "summary": summary,
        "symbol": clean_symbol,
        "pnl_headers": pnl_headers,
        "pnl": pnl,
        "bs_headers": bs_headers,
        "bs": bs,
        "cf_headers": cf_headers,
        "cf": cf,
        "ratio_headers": ratio_headers,
        "ratios": ratios,
        "peer_headers": peer_headers,
        "peers": peer_rows,
        "share_headers": share_headers,
        "shareholding": shareholding,
        "candles": json.dumps(candles),
        "period": period,
        "context_periods": period_options,
    }

    return render(request, "stock_detail.html", context)



def register(request):
    if request.method == "POST":
        request.session["user_data"] = {
            "first_name": request.POST.get("first_name"),
            "last_name": request.POST.get("last_name"),
            "email": request.POST.get("email"),
            "phone": request.POST.get("phone"),
            "password": request.POST.get("password"),
            "dob": request.POST.get("dob"),
            "gender": request.POST.get("gender"),
            "address": request.POST.get("address"),
        }

        client = razorpay.Client(
            auth=(settings.RAZORPAY_KEY_ID, settings.RAZORPAY_KEY_SECRET)
        )

        order = client.order.create({
            "amount": 100000,
            "currency": "INR",
            "payment_capture": "1",
        })

        request.session["order_id"] = order["id"]

        context = {
            "razorpay_key": settings.RAZORPAY_KEY_ID,
            "order_id": order["id"],
            "amount": 100000,
        }

        return render(request, "payment.html", context)

    return render(request, "register.html")



def payment_success(request):
    if request.method != "POST":
        return redirect("register")

    order_id = request.POST.get("razorpay_order_id")
    payment_id = request.POST.get("razorpay_payment_id")
    signature = request.POST.get("razorpay_signature")
    expected_order_id = request.session.get("order_id")
    data = request.session.get("user_data")

    if not data:
        return redirect("register")

    if expected_order_id and order_id != expected_order_id:
        return render(request, "payment.html", {
            "error": "Order mismatch. Please try the payment again."
        })

    if signature and settings.RAZORPAY_KEY_ID and settings.RAZORPAY_KEY_SECRET:
        client = razorpay.Client(
            auth=(settings.RAZORPAY_KEY_ID, settings.RAZORPAY_KEY_SECRET)
        )
        try:
            client.utility.verify_payment_signature({
                "razorpay_order_id": order_id,
                "razorpay_payment_id": payment_id,
                "razorpay_signature": signature,
            })
        except razorpay.errors.SignatureVerificationError:
            return render(request, "payment.html", {
                "error": "Payment signature verification failed."
            })

    user = User.objects.create(
        first_name=data["first_name"],
        last_name=data["last_name"],
        email=data["email"],
        phone=data["phone"],
        password=hashlib.md5(data["password"].encode()).hexdigest(),
        dob=data["dob"],
        gender=data["gender"],
        address=data["address"],
    )

    Payment.objects.create(
        user=user,
        razorpay_order_id=order_id,
        razorpay_payment_id=payment_id,
        amount=1000,
    )

    request.session.pop("user_data", None)
    request.session.pop("order_id", None)

    return redirect("login")



def login_view(request):
    if request.method == "POST":
        email = request.POST.get("email")
        password = request.POST.get("password")
        remember_me = request.POST.get("remember_me")

        hashed_password = hashlib.md5(password.encode()).hexdigest()

        try:
            user = User.objects.get(email=email, password=hashed_password)

            request.session["user_id"] = user.id
            request.session["user_name"] = user.first_name

            if not remember_me:
                request.session.set_expiry(0)
            else:
                request.session.set_expiry(1209600)

            return redirect("welcome")

        except User.DoesNotExist:
            return render(request, "login.html", {"error": "Invalid Email or Password"})

    return render(request, "login.html")



def forgot_password(request):
    if request.method == "POST":
        email = request.POST.get("email")

        try:
            user = User.objects.get(email=email)

            token = str(uuid.uuid4())
            user.reset_token = token
            user.token_created_at = timezone.now()
            user.save()

            reset_link = request.build_absolute_uri(
                reverse("email_reset", args=[token])
            )

            send_mail(
                "Password Reset Request",
                f"Click the link below to reset your password:\n\n{reset_link}",
                settings.EMAIL_HOST_USER,
                [email],
                fail_silently=False,
            )

            return render(request, "forgot.html", {"message": "Reset link sent to your email."})

        except User.DoesNotExist:
            return render(request, "forgot.html", {"error": "Email not registered."})

    return render(request, "forgot.html")



def reset_password(request):
    context = {}

    if request.session.get("reset_verified"):
        context["verified"] = True
        context["email"] = request.session.get("reset_email")

    if request.method == "POST":
        if not request.session.get("reset_verified"):
            email = request.POST.get("email")

            try:
                User.objects.get(email=email)

                request.session["reset_verified"] = True
                request.session["reset_email"] = email

                context["verified"] = True
                context["email"] = email

            except User.DoesNotExist:
                context["error"] = "Invalid Email"
        else:
            new_password = request.POST.get("new_password")
            email = request.session.get("reset_email")
            user = User.objects.get(email=email)

            hashed_password = hashlib.md5(new_password.encode()).hexdigest()
            user.password = hashed_password
            user.save()

            request.session.pop("reset_verified", None)
            request.session.pop("reset_email", None)

            return redirect("login")

    return render(request, "reset.html", context)



def email_reset(request, token):
    try:
        user = User.objects.get(reset_token=token)

        if timezone.now() > user.token_created_at + timezone.timedelta(minutes=10):
            return render(request, "reset.html", {"error": "Reset link expired."})

        if request.method == "POST":
            new_password = request.POST.get("new_password")

            hashed_password = hashlib.md5(new_password.encode()).hexdigest()
            user.password = hashed_password
            user.reset_token = None
            user.token_created_at = None
            user.save()

            return redirect("login")

        return render(request, "reset.html")

    except User.DoesNotExist:
        return render(request, "reset.html", {"error": "Invalid reset link."})



def _get_watchlist_credentials():
    return (
        getattr(settings, "SCREENER_EMAIL", "xyzpqr456@gmail.com"),
        getattr(settings, "SCREENER_PASSWORD", "Stock@123"),
        getattr(settings, "SCREENER_WATCHLIST_URL", "https://www.screener.in/watchlist/10077488/"),
    )



def _load_watchlist_data(request, force_refresh=False):
    headers = request.session.get("watchlist_headers")
    rows = request.session.get("watchlist_rows")

    if headers and rows and not force_refresh:
        return headers, rows

    screener_email, screener_password, watchlist_url = _get_watchlist_credentials()
    headers, rows = scrape_watchlist_with_selenium(
        screener_email,
        screener_password,
        watchlist_url,
    )

    request.session["watchlist_headers"] = headers
    request.session["watchlist_rows"] = rows
    return headers, rows



def _build_watchlist_pdf(headers, rows, generated_for):
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(legal),
        rightMargin=20,
        leftMargin=20,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Stock Watchlist Dashboard Report", styles["Title"]),
        Spacer(1, 8),
        Paragraph(f"Generated for: {generated_for}", styles["Normal"]),
        Paragraph(
            f"Generated on: {timezone.localtime(timezone.now()).strftime('%d %b %Y, %I:%M %p')}",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]

    pdf_rows = build_watchlist_report_rows(rows)
    table_data = [headers or WATCHLIST_FULL_HEADERS] + pdf_rows

    pdf_table = Table(table_data, repeatRows=1)
    pdf_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5a1fd1")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#eef2ff")]),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
    ]))

    elements.append(pdf_table)
    document.build(elements)

    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes



def welcome(request):
    if "user_id" not in request.session:
        return redirect("login")

    user = User.objects.get(id=request.session["user_id"])

    try:
        scraped_headers, rows = _load_watchlist_data(request)
    except Exception as exc:
        scraped_headers, rows = [], []
        messages.error(request, f"Unable to load watchlist right now: {exc}")

    stocks = Stock.objects.all().values("name", "slug")

    display_headers = WATCHLIST_FULL_HEADERS
    base_headers = scraped_headers or WATCHLIST_BASE_HEADERS

    return render(request, "welcome.html", {
        "user": user,
        "watchlist_headers": base_headers,
        "watchlist_rows": rows,
        "display_headers": display_headers,
        "stocks": list(stocks),
    })



def send_watchlist_email(request):
    if request.method != "POST":
        return redirect("welcome")

    if "user_id" not in request.session:
        return redirect("login")

    user = User.objects.get(id=request.session["user_id"])

    try:
        base_headers, rows = _load_watchlist_data(request)

        if not rows:
            base_headers, rows = _load_watchlist_data(request, force_refresh=True)

        if not rows:
            messages.error(request, "No watchlist data is available to email.")
            return redirect("welcome")

        pdf_headers = WATCHLIST_FULL_HEADERS
        pdf_bytes = _build_watchlist_pdf(pdf_headers, rows, user.email)

        email = EmailMessage(
            subject="Your Stock Watchlist PDF Report",
            body=(
                f"Hello {user.first_name},\n\n"
                "Your latest stock watchlist report is attached as a PDF file.\n\n"
                "Regards,\n"
                "Stock Analyzer"
            ),
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[user.email],
        )
        email.attach("stock_watchlist_report.pdf", pdf_bytes, "application/pdf")
        email.send(fail_silently=False)

        request.session["watchlist_headers"] = base_headers
        request.session["watchlist_rows"] = rows
        messages.success(request, f"Watchlist PDF sent successfully to {user.email}.")
    except Exception as exc:
        messages.error(request, f"Could not send the email update: {exc}")

    return redirect("welcome")



def logout_view(request):
    request.session.flush()
    return redirect("login")



def export_excel(request):
    rows = request.session.get("watchlist_rows", [])

    if not rows:
        return HttpResponse("No data available. Please load dashboard first.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Watchlist"

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    headers = [
        "S.No.", "", "Name", "CMP(Rs.)", "P/E Ratio", "Price to Book Value",
        "PEG Ratio", "Market Cap/Sales", "Sales Growth 5Y(%)",
        "Profit Growth 5Y(%)", "ROE(%)", "ROCE(%)",
        "Asset Turnover Ratio", "OPM(%)", "NPM(%)", "Debt/Equ",
        "Interest Coverage Ratio", "Inventory Turnover Ratio",
        "Current Ratio"
    ]

    ws.append(headers)

    for r, row in enumerate(rows, start=2):
        normalized_row = list(row[:19]) + [""] * max(0, 19 - len(row))
        ws.append(normalized_row)

        pe = safe_float(normalized_row[4])
        cmpbv = safe_float(normalized_row[5])
        peg = safe_float(normalized_row[6])
        mcapsales = safe_float(normalized_row[7])
        sales = safe_float(normalized_row[8])
        profit = safe_float(normalized_row[9])
        roe = safe_float(normalized_row[10])
        roce = safe_float(normalized_row[11])
        asset = safe_float(normalized_row[12])
        opm = safe_float(normalized_row[13])
        npm = safe_float(normalized_row[14])
        debt = safe_float(normalized_row[15])
        interest = safe_float(normalized_row[16])
        invent = safe_float(normalized_row[17])
        current = safe_float(normalized_row[18])

        rules = [
            (5, pe < 30),
            (6, cmpbv < 3),
            (7, peg < 1),
            (8, mcapsales < 3),
            (9, sales > 10),
            (10, profit > 15),
            (11, roe > 15),
            (12, roce > 15),
            (13, asset > 1),
            (14, opm > 15),
            (15, npm > 10),
            (16, debt <= 0.2),
            (17, interest >= 20),
            (18, invent > 4),
            (19, 1.5 <= current <= 2),
        ]

        for col, cond in rules:
            cell = ws.cell(row=r, column=col)
            cell.fill = green if cond else red

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=stock_watchlist.xlsx"

    wb.save(response)
    return response



def about(request):
    return render(request, "about.html")



def privacy(request):
    return render(request, "privacy.html")
